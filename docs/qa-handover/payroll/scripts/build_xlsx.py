import json, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from collections import Counter
rows=json.load(open('../logs/reconciliation.json'))
RED=PatternFill('solid',fgColor='F8CBAD'); YEL=PatternFill('solid',fgColor='FFF2CC'); GRN=PatternFill('solid',fgColor='E2EFDA'); HDR=PatternFill('solid',fgColor='1F4E78')
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Summary'
def hdr(s,cols):
    s.append(cols)
    for c in s[s.max_row]: c.font=Font(bold=True,color='FFFFFF'); c.fill=HDR; c.alignment=Alignment(wrap_text=True,vertical='top')
cnt=Counter(r['status'] for r in rows); cm=Counter(r['status'] for r in rows if r['mock'])
fc=Counter(f for r in rows for f in r['mismatch_fields'] if '~' not in f)
pf_only=sum(1 for r in rows if r['status']=='MISMATCH' and set(f for f in r['mismatch_fields'] if '~' not in f)<= {'pf','ded','net'})
ws.append(['Rashmi HRMS payroll reconciliation - June 2026 (runs 4 RML, 5 RGH), DB hrms, backend :5100']); ws['A1'].font=Font(bold=True,size=13)
ws.append(['Generated 2026-10-18 IST. Expected values computed independently (scripts/reconcile.py); app values from pay.payroll_items/lines after final compute.'])
ws.append([])
hdr(ws,['Metric','All items (44)','Mock cohort (33)'])
for k in ('MATCH','MISMATCH'): ws.append([k,cnt.get(k,0),cm.get(k,0)])
ws.append(['Mismatches explained only by PF ceiling (PF/ded/net only)',pf_only,sum(1 for r in rows if r['mock'] and r['status']=='MISMATCH' and set(f for f in r['mismatch_fields'] if '~' not in f)<={'pf','ded','net'})])
ws.append([]); hdr(ws,['Field','Items with |diff| >= Rs 1',''])
for k,v in fc.most_common(): ws.append([k,v])
ws.append([]); hdr(ws,['Totals','Expected','App'])
for f in ('gross','pf','esi','pt','tds','other','ded','net'):
    ws.append([f,sum(r['exp_'+f] for r in rows),round(sum(r['app_'+f] for r in rows),2)])
ws.column_dimensions['A'].width=60; ws.column_dimensions['B'].width=18; ws.column_dimensions['C'].width=18

s=wb.create_sheet('Reconciliation')
cols=['run','ecode','name','mock','case','status','mismatch_fields','exp_days','app_days','d_days','days_note','exp_gross','app_gross','d_gross','exp_pf','app_pf','d_pf','exp_esi','app_esi','d_esi','exp_pt','app_pt','d_pt','exp_tds','app_tds','d_tds','exp_other','app_other','d_other','exp_ded','app_ded','d_ded','exp_net','app_net','d_net','payment_status','first_compute_net']
hdr(s,cols)
for r in rows:
    s.append([', '.join(r[c]) if c=='mismatch_fields' else r[c] for c in cols])
    row=s[s.max_row]
    row[cols.index('status')].fill = GRN if r['status']=='MATCH' else (YEL if r['status'].startswith('ROUND') else RED)
    for f in ('days','gross','pf','esi','pt','tds','other','ded','net'):
        d=r['d_'+f]
        if abs(d)>=1:
            for c in ('exp_','app_','d_'): row[cols.index(c+f)].fill=RED
        elif d: row[cols.index('d_'+f)].fill=YEL
s.freeze_panes='D2'
for i,w in enumerate([5,11,22,6,45,12,22]): s.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w

s=wb.create_sheet('Sample RML033625')
smp=next(r for r in rows if r['ecode']=='RML033625')
hdr(s,['Line','Register (ground truth)','App run 4','Diff (app-register)','Explanation'])
reg=[('BASIC',32286),('HRA',16143),('EDUCATION_ALLOW',200),('MEDICAL_ALLOW',1250),('STATUTORY_BONUS',2689),('SPECIAL_ALLOW',12005),('GROSS',64573),('PF_EE',3874),('ESI_EE',0),('PT',200),('TDS',0),('CANTEEN',1375),('TOTAL DEDUCTIONS',5449),('NET',59124)]
L=smp['app_lines']
appv={'GROSS':smp['app_gross'],'TOTAL DEDUCTIONS':smp['app_ded'],'NET':smp['app_net']}
expl={'PF_EE':'Engine caps PF wages at Rs15,000 (12% x 15000 = 1800). Register uses 12% of full basic 32286 = 3874. compute-engine.ts:149 & 284-287 default ceiling; run-compute.service.ts:190-202 never passes pfWageCeilingPaise / company setting.',
      'TOTAL DEDUCTIONS':'= PF difference (-2074)','NET':'= PF difference (+2074); all other lines identical'}
for k,v in reg:
    a=appv.get(k, L.get(k,0)); d=round(a-v,2)
    s.append([k,v,a,d,expl.get(k,'match' if d==0 else '')])
    if d: 
        for c in s[s.max_row]: c.fill=RED
s.column_dimensions['A'].width=20; s.column_dimensions['E'].width=110

s=wb.create_sheet('Employer contributions')
hdr(s,['run','ecode','exp PF_ER (3.67%)','exp EPS (8.33% capped)','exp EDLI','exp ESI_ER','app employer lines'])
for r in rows:
    e=r['exp_er']; s.append([r['run'],r['ecode'],e.get('EMP_PF_ER',0),e.get('EMP_EPS',0),e.get('EMP_EDLI',0),e.get('EMP_ESI_ER',0),', '.join(k for k in r['app_lines'] if k in ('PF_ER','EPS_ER','EDLI','PF_ADMIN','ESI_ER')) or 'NONE'])
    if s[s.max_row][6].value=='NONE': s[s.max_row][6].fill=RED

s=wb.create_sheet('Line detail')
hdr(s,['run','ecode','component','expected','app','diff'])
for r in rows:
    keys=sorted(set(k for k in r['exp_lines'] if not k.startswith('EMP_'))|set(r['app_lines']))
    for k in keys:
        e=r['exp_lines'].get(k,0); a=r['app_lines'].get(k,0); d=round(a-e,2)
        s.append([r['run'],r['ecode'],k,e,a,d])
        if abs(d)>=1: s[s.max_row][5].fill=RED
        elif d: s[s.max_row][5].fill=YEL

s=wb.create_sheet('Conventions')
for t in ['Days in month 30 (calendar). Earnings = monthly component x payable/30, rounded to rupee (half-up). Statutory bonus prorated (Payment of Bonus Act: on wages earned).',
 'Inputs (pay.inputs) added as entered. App loan (loan 4, 60000/12) -> LOAN_EMI 5000 expected for RML033917.',
 'OT RGH033901: BASIC/208 x 2 x 4h = 1058 (Factories Act double rate). App: approved OT made compute throw; entry set rejected by SQL to let run compute (documented workaround).',
 'PF_EE = 12% of earned BASIC, ACTUAL (no 15,000 ceiling) per sample register. Apprentice: no PF/ESI/bonus.',
 'ESI_EE = ceil(0.75% x earned gross) only if employee is ESIC-applicable AND fixed monthly gross <= 21,000.',
 'PT: WB slab on earned gross (<=10k 0, <=15k 110, <=25k 130, <=40k 150, >40k 200). Odisha (RML033923): annual >3L -> 200/month (assumption).',
 'TDS: new regime std ded 75k, slabs 4/8/12/16/20/24L, 87A rebate to 12L with marginal relief, surcharge 10% >50L with marginal relief, cess 4%; old regime std ded 50k, 80C 1.5L, PT deductible; annual / 12. Only for full-month employees.',
 'Maternity (RML033904): ML 08-30 Jun paid -> 30 payable days. Process holds (RGH033908, RML033922 after workaround) -> zero item expected.',
 'Match tolerance: |diff| < Rs 1 = rounding (yellow); >= Rs 1 = mismatch (red).']:
    s.append([t])
s.column_dimensions['A'].width=160
wb.save('../reconciliation.xlsx'); print('saved')
